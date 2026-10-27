import io

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)
MD = "# Audit\n\n## Findings\n\n- **Missing**: breach notice\n1. Add clause\n\n| Status | Count |\n|---|---|\n| Covered | 9 |\n\nहिंदी सारांश"


def _export(fmt, **kw):
    return client.post("/api/export", json={"markdown": MD, "title": "Zomato audit", "format": fmt, **kw})


def test_pdf_is_a_real_multi_element_pdf():
    import pymupdf
    r = _export("pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert 'filename="zomato-audit.pdf"' in r.headers["content-disposition"]
    doc = pymupdf.open("pdf", r.content)
    text = doc[0].get_text()
    assert "Zomato audit" in text and "Findings" in text and "Page 1 of 1" in text
    assert doc.metadata["title"] == "Zomato audit"


def test_docx_uses_real_word_structure():
    from docx import Document
    r = _export("docx")
    assert r.status_code == 200 and r.headers["content-type"].endswith("wordprocessingml.document")
    d = Document(io.BytesIO(r.content))
    styles = {p.style.name for p in d.paragraphs}
    assert {"Title", "Heading 2", "List Bullet", "List Number"} <= styles
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "9"
    assert any(r.bold and r.text == "Missing" for p in d.paragraphs for r in p.runs)


def test_markdown_gets_title_block():
    r = _export("md")
    assert r.status_code == 200 and r.text.startswith("# Zomato audit") and "## Findings" in r.text


def test_rejects_bad_requests():
    assert client.post("/api/export", json={"markdown": "  ", "title": "x", "format": "pdf"}).status_code == 400
    assert client.post("/api/export", json={"markdown": "x", "title": "x", "format": "exe"}).status_code == 422

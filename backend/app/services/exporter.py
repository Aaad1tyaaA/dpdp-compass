"""Export Markdown content (audit reports, policy drafts, chat answers) as Markdown, PDF or Word.

Markdown is the single source of truth; PDF and DOCX are rendered from it so all three downloads
say exactly the same thing. Both rendered formats use A4, a title block, a heading hierarchy,
real lists/tables, and "Page X of Y" footers.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

ACCENT = (79, 70, 229)  # indigo-600, matches the web UI accent
MUTED = (100, 116, 139)  # slate-500
INK = (15, 23, 42)  # slate-900


# ── helpers ───────────────────────────────────────────────────────────────────
def _stamp() -> str:
    return datetime.now().strftime("%d %B %Y")


def _strip_leading_title(markdown: str, title: str) -> str:
    """Drop a leading '# Title' when it repeats the document title (the title block already shows it)."""
    lines = markdown.lstrip().splitlines()
    if lines and lines[0].startswith("# ") and lines[0][2:].strip().lower() == title.strip().lower():
        return "\n".join(lines[1:]).lstrip()
    return markdown


# ── PDF ───────────────────────────────────────────────────────────────────────
_WIN_FONTS = Path("C:/Windows/Fonts")


def _pdf_css() -> str:
    """Segoe UI when available (Windows), with Nirmala UI for Hindi/Marathi; MuPDF's Helvetica otherwise."""
    faces = ""
    family = "sans-serif"
    if (_WIN_FONTS / "segoeui.ttf").exists():
        faces = (
            "@font-face{font-family:Body;src:url(segoeui.ttf);}"
            "@font-face{font-family:Body;src:url(segoeuib.ttf);font-weight:bold;}"
            "@font-face{font-family:Body;src:url(segoeuii.ttf);font-style:italic;}"
        )
        family = "Body"
        if (_WIN_FONTS / "Nirmala.ttc").exists() or (_WIN_FONTS / "Nirmala.ttf").exists():
            src = "Nirmala.ttc" if (_WIN_FONTS / "Nirmala.ttc").exists() else "Nirmala.ttf"
            faces += f"@font-face{{font-family:Deva;src:url({src});}}@font-face{{font-family:Deva;src:url({src});font-weight:bold;}}"
            family = "Body, Deva"
    ink = "rgb(%d,%d,%d)" % INK
    accent = "rgb(%d,%d,%d)" % ACCENT
    muted = "rgb(%d,%d,%d)" % MUTED
    return faces + f"""
      * {{ font-family: {family}; }}
      body {{ font-size: 10.5pt; line-height: 1.5; color: {ink}; }}
      .doc-title {{ font-size: 21pt; font-weight: bold; color: {ink}; margin: 0 0 2pt 0; text-align: left; line-height: 1.2; }}
      .doc-sub {{ font-size: 9.5pt; color: {muted}; margin: 0 0 14pt 0; padding-bottom: 8pt; border-bottom: 1.2pt solid {accent}; }}
      h1, h2, h3, h4 {{ text-align: left; }}
      h1 {{ font-size: 15pt; margin: 16pt 0 6pt 0; color: {ink}; }}
      h2 {{ font-size: 12.5pt; margin: 14pt 0 5pt 0; color: {accent}; }}
      h3, h4 {{ font-size: 11pt; margin: 10pt 0 4pt 0; color: {ink}; }}
      p {{ margin: 0 0 7pt 0; text-align: justify; }}
      ul, ol {{ margin: 0 0 8pt 0; padding-left: 16pt; }}
      li {{ margin: 0 0 3pt 0; }}
      strong {{ font-weight: bold; }}
      code {{ font-family: monospace; font-size: 9.5pt; color: rgb(51,65,85); }}
      hr {{ border: 0; border-top: 0.6pt solid rgb(226,232,240); margin: 10pt 0; }}
      table {{ border-collapse: collapse; width: 100%; margin: 4pt 0 10pt 0; }}
      th {{ font-weight: bold; text-align: left; color: {accent}; border-bottom: 1.2pt solid {accent}; }}
      /* no background fills: MuPDF's story layout repaints cell backgrounds on every later page */
      th, td {{ border: 0.6pt solid rgb(203,213,225); padding: 4pt 6pt; font-size: 9.5pt; vertical-align: top; }}
      blockquote {{ margin: 0 0 8pt 0; padding-left: 8pt; border-left: 2pt solid {accent}; color: {muted}; }}
    """


def to_pdf(markdown: str, title: str, subtitle: str = "") -> bytes:
    import markdown as md
    import pymupdf

    body = md.markdown(_strip_leading_title(markdown, title), extensions=["tables", "sane_lists"])
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    sub = esc(subtitle or f"DPDP Compass · Generated {_stamp()}")
    html = f'<p class="doc-title">{esc(title)}</p><p class="doc-sub">{sub}</p>{body}'

    archive = pymupdf.Archive(str(_WIN_FONTS)) if _WIN_FONTS.exists() else None
    story = pymupdf.Story(html=html, user_css=_pdf_css(), archive=archive)
    page_rect = pymupdf.paper_rect("a4")
    content = page_rect + (56, 64, -56, -60)  # ~2cm side margins, room for header/footer

    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Running header + "Page X of Y" footer (page count is only known after layout).
    doc = pymupdf.open("pdf", buf.getvalue())
    grey = tuple(c / 255 for c in MUTED)
    n = doc.page_count
    for i, page in enumerate(doc, start=1):
        w, h = page.rect.width, page.rect.height
        if i > 1:
            page.insert_text((56, 38), title[:90], fontsize=8, color=grey)
        page.draw_line((56, h - 44), (w - 56, h - 44), color=(0.89, 0.91, 0.94), width=0.6)
        page.insert_text((56, h - 30), "DPDP Compass · Educational guidance, not legal advice", fontsize=7.5, color=grey)
        label = f"Page {i} of {n}"
        page.insert_text((w - 56 - pymupdf.get_text_length(label, fontsize=7.5), h - 30), label, fontsize=7.5, color=grey)
    doc.set_metadata({"title": title, "author": "DPDP Compass", "subject": subtitle or title, "creator": "DPDP Compass"})
    doc.subset_fonts()  # embed only the glyphs used (whole Segoe/Nirmala files made a 1.5 MB PDF)
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


# ── DOCX ──────────────────────────────────────────────────────────────────────
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|(?<!\w)_[^_\s][^_]*_(?!\w)|`[^`]+`)")


def _add_runs(paragraph, text: str) -> None:
    from docx.shared import Pt, RGBColor
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name, r.font.size = "Consolas", Pt(9.5)
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif len(part) > 2 and part[0] in "*_" and part[-1] == part[0]:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _style_table(table) -> None:
    """Clean report table: tinted bold header row, thin light borders, no banding."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    table.style = "Table Grid"
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "CBD5E1")
        borders.append(el)
    table._tbl.tblPr.append(borders)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    if ri == 0:
                        run.bold, run.font.color.rgb = True, RGBColor(*ACCENT)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear"); shade.set(qn("w:color"), "auto"); shade.set(qn("w:fill"), "EEF2FF")
                cell._tc.get_or_add_tcPr().append(shade)


def _page_number_field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run = paragraph.add_run()
    run.font.size, run.font.color.rgb = Pt(8), RGBColor(*MUTED)  # same size as the footer text
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


def _style_document(doc) -> None:
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Segoe UI", Pt(10.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    base.element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")  # Hindi/Marathi (complex script) text
    base.paragraph_format.space_after = Pt(6)
    base.paragraph_format.line_spacing = 1.15
    ink, accent = RGBColor(*INK), RGBColor(*ACCENT)
    for name, size, color in (("Title", 22, ink), ("Heading 1", 15, ink), ("Heading 2", 12.5, accent), ("Heading 3", 11, ink)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Segoe UI", Pt(size), True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
        st.paragraph_format.space_before = Pt(0 if name == "Title" else 12)
        st.paragraph_format.space_after = Pt(4)


def to_docx(markdown: str, title: str, subtitle: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    _style_document(doc)
    doc.core_properties.title, doc.core_properties.author = title, "DPDP Compass"

    doc.add_paragraph(title, style="Title")
    sub = doc.add_paragraph()
    r = sub.add_run(subtitle or f"DPDP Compass · Generated {_stamp()}")
    r.font.size, r.font.color.rgb = Pt(9.5), RGBColor(*MUTED)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("DPDP Compass · Educational guidance, not legal advice    Page ")
    fr.font.size, fr.font.color.rgb = Pt(8), RGBColor(*MUTED)
    _page_number_field(footer, "PAGE")
    of = footer.add_run(" of ")
    of.font.size, of.font.color.rgb = Pt(8), RGBColor(*MUTED)
    _page_number_field(footer, "NUMPAGES")

    lines = _strip_leading_title(markdown, title).splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{3,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    _add_runs(table.cell(ri, ci).paragraphs[0], cell)
            _style_table(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)  # breathing room after the table
            continue
        h = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        answer_heading = re.match(r"^\*\*(\d+\.\s.+)\*\*$", stripped)  # model answers: "**1. Short Answer**"
        if h:
            level = min(len(h.group(1)), 3)
            doc.add_paragraph(h.group(2).strip("# ").replace("**", ""), style=f"Heading {level}")
        elif answer_heading:
            doc.add_paragraph(answer_heading.group(1), style="Heading 3")
        elif re.match(r"^-{3,}$", stripped):
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif m := re.match(r"^\s*[-*•]\s+(.*)$", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)$", line):
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            _add_runs(p, stripped.lstrip("> "))
            for run in p.runs:
                run.italic, run.font.color.rgb = True, RGBColor(*MUTED)
        else:
            _add_runs(doc.add_paragraph(), stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── entry point ───────────────────────────────────────────────────────────────
FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
}


def _normalise(markdown: str) -> str:
    """Model-style source markers (fullwidth brackets) read badly in documents; use plain ones."""
    return re.sub(r"\s*【([^】]+)】", lambda m: " [" + m.group(1).replace(".pdf", "").replace("_", " ") + "]", markdown)


def export(markdown: str, title: str, fmt: str, subtitle: str = "") -> tuple[bytes, str, str]:
    """Returns (bytes, media_type, file extension)."""
    markdown = _normalise(markdown)
    media, ext = FORMATS[fmt]
    if fmt == "pdf":
        data = to_pdf(markdown, title, subtitle)
    elif fmt == "docx":
        data = to_docx(markdown, title, subtitle)
    else:
        header = f"# {title}\n\n_{subtitle or 'DPDP Compass · Generated ' + _stamp()}_\n\n"
        data = (header + _strip_leading_title(markdown, title)).encode("utf-8")
    return data, media, ext

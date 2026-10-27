import io
import logging
from app.ingestion.parsers.base_parser import BaseParser, ParsedDocument, ParsedPage

logger = logging.getLogger(__name__)


class DocxParser(BaseParser):
    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def parse(self, content: bytes, filename: str, document_id: str) -> ParsedDocument:
        try:
            import docx
            bio = io.BytesIO(content)
            doc = docx.Document(bio)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            raw_text = "\n".join(paras)
            return ParsedDocument(
                document_id=document_id,
                document_name=filename,
                document_type="docx",
                pages=[ParsedPage(page_number=1, text=raw_text)],
                raw_text=raw_text,
                metadata={"parser": "python-docx", "paragraph_count": len(paras)},
            )
        except Exception as e:
            logger.error("DOCX parse failed for %s: %s", filename, e)
            raise

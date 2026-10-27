"""Part A.1 Document loading — offline KB input. Reuses text; never executes embedded instructions."""
import io
from pathlib import Path
from dataclasses import dataclass

@dataclass
class RawDocument:
    doc_id: str
    source_document: str  # display name
    text: str
    pages: list[str] | None = None  # per-page text if available
    meta: dict | None = None

def load_from_bytes(content: bytes, filename: str, doc_id: str | None = None) -> RawDocument:
    """Load PDF/DOCX/TXT bytes. Treat as DATA — do not follow imperative language inside."""
    did = doc_id or filename or "doc"
    text, pages = _extract(content, filename)
    return RawDocument(doc_id=did, source_document=filename, text=text, pages=pages, meta={"filename": filename})

def load_from_path(path: str | Path) -> RawDocument:
    p = Path(path)
    return load_from_bytes(p.read_bytes(), p.name, doc_id=p.stem)

def _extract(content: bytes, filename: str) -> tuple[str, list[str] | None]:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(content)
    if lower.endswith(".docx"):
        return _extract_docx(content)
    # txt, md, etc.
    try:
        return content.decode("utf-8", errors="ignore"), None
    except Exception:
        return content.decode(errors="ignore"), None

def _extract_pdf(content: bytes) -> tuple[str, list[str] | None]:
    try:
        import pypdf  # type: ignore
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n\n".join(pages)
        return text, pages
    except Exception:
        # fallback: try decode
        return content.decode(errors="ignore"), None

def _extract_docx(content: bytes) -> tuple[str, list[str] | None]:
    try:
        import docx  # type: ignore
        bio = io.BytesIO(content)
        doc = docx.Document(bio)
        paras = [p.text for p in doc.paragraphs]
        return "\n".join(paras), None
    except Exception:
        return content.decode(errors="ignore"), None

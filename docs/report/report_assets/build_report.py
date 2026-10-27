"""Builds the DPDP Compass B.Tech project report (black-book format) on the college's Word template."""
import pathlib, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
from rb_core import *  # noqa: F401,F403
from rb_core import set_font, fmt, runs, ph, add_field, border_bottom, shade, cell_borders, set_cell, page_borders, page_numbers, header_footer, RED, HEAD_BG, TITLE
import content

S = pathlib.Path(__file__).parent
TPL = S / "tpl" / "1__Title page, Certificate, Acknowledgement.docx"
OUT = pathlib.Path(sys.argv[1])
C, L, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY
FULL = "Multi-Agent Retrieval-Augmented Generation System for DPDP Act Compliance Assessment and Citizen Rights Awareness"
TEAM = [("Raj Ashok Nikam", 6, "EN23225464"), ("Aditya Sachin Patil", 7, "EN23238761"), ("Varad Rahul Mhetar", 8, "EN23176065"),
        ("Shubham Sanjaykumar Rajurkar", 9, "EN23149824"), ("Shubhdeep Sunil Jasuja", 10, "EN23222171")]
GUIDE, COORD, HOD, PRIN = "Mrs. D. R. Patil", "Dr. A. S. Yadav", "Ms. R. J. Dhanal", "Prof. Dr. S. S. Pawar"

doc = Document(TPL)
doc.settings.odd_and_even_pages_header_footer = False
body = doc.element.body
els = list(body.iterchildren())
P = lambda i: __import__("docx").text.paragraph.Paragraph(els[i], doc)
T = lambda i: __import__("docx").table.Table(els[i], doc)


def retext(p, text, size=None, bold=None, italic=None):
    first = p.runs[0] if p.runs else None
    sz = size or (first.font.size.pt if first is not None and first.font.size else None); bd = first.bold if bold is None and first is not None else bold
    for r in list(p.runs): r._element.getparent().remove(r._element)
    runs(p, text, sz or 12, bool(bd)); return p


# ── front matter from the template ─────────────────────────────────────────
retext(P(7), FULL, 16, True)
_g = doc.add_paragraph(); fmt(_g, C, 2, 6, 1.15); runs(_g, "Group ID: **G45**", 13); els[8].addnext(_g._p)
for i in (16, 13):  # two empty spacer lines on the title page make room for the Group ID
    if not "".join(els[i].itertext()).strip(): els[i].getparent().remove(els[i])
t = T(9)
for i, (n, roll, en) in enumerate(TEAM, 1):
    for j, v in enumerate((str(i), n, str(roll), en)): set_cell(t.rows[i].cells[j], v, 11, align=C if j != 1 else L)
retext(P(12), GUIDE, None, True)
g = T(15).rows[0].cells
set_cell(g[0], "**Group number:** G45", 11); set_cell(g[1], "**Class:** B.Tech (CSE)", 11); set_cell(g[2], "**Division:** D", 11)
cert = P(26)
for r in list(cert.runs): r._element.getparent().remove(r._element)
runs(cert, f"This is to certify that the project group **‘G45’**, consisting of the following members, has satisfactorily completed the Project work entitled "
     f"**“{FULL}”** at B. Tech (CSE), Semester – VII as prescribed in the syllabus for the academic year 2026–2027.", 12)
t = T(27)
for i, (n, roll, en) in enumerate(TEAM, 1):
    vals = (str(i), n, str(roll), en, "[[Exam No.]]")
    for j, v in enumerate(vals): set_cell(t.rows[i].cells[j], v, 8.5 if j == 4 else 10.5, align=C if j != 1 else L)
sig = T(29)
set_cell(sig.rows[1].cells[0], f"**{GUIDE}**", 11, align=C); set_cell(sig.rows[1].cells[2], f"**{COORD}**", 11, align=C)
set_cell(sig.rows[7].cells[1], "[[Name of External Examiner]]", 11, align=C)
sig.rows[7].cells[1].add_paragraph(); p = sig.rows[7].cells[1].paragraphs[-1]; fmt(p, C, 0, 0, 1.0); runs(p, "External Examiner", 11)
retext(P(31), "**Date:** [[DD/MM/YYYY]]          **Place:** Kolhapur", 12); els[32].getparent().remove(els[32])
# keep the certificate on one page: drop the signature table's blank first and seventh rows and two empty spacer paragraphs
for r in (sig.rows[6], sig.rows[0]): r._tr.getparent().remove(r._tr)
for i in (33, 30):
    if not "".join(els[i].itertext()).strip(): els[i].getparent().remove(els[i])
for c in sig.rows[0].cells: c.paragraphs[0].paragraph_format.space_before = Pt(8)
# the certificate fills its page, so the template's separate page-break paragraph would land on a blank page:
# end the Date/Place line with the break instead and drop that paragraph
P(31).add_run().add_break(WD_BREAK.PAGE)
els[34].getparent().remove(els[34])
# the acknowledgement title is a heading in the template; keep it out of the table of contents
ack_title = P(35); ack_title.style = doc.styles["Normal"]
fmt(ack_title, C, 0, 12, 1.15); retext(ack_title, "**Acknowledgement**", 16, True)
ack1 = P(38)
for r in list(ack1.runs): r._element.getparent().remove(r._element)
runs(ack1, f"We would like to express our profound gratitude to our internal guide of the project, **{GUIDE}**, [[Designation]], Department of Computer Science "
     "and Engineering, D. Y. Patil College of Engineering and Technology, Kolhapur, for the continuous support of our Project Work, for her patience, "
     "motivation, encouragement, and immense knowledge.", 12)
ack2 = P(39)
for r in list(ack2.runs): r._element.getparent().remove(r._element)
runs(ack2, f"We would like to express our gratitude and deep regards to the Department Project Coordinator, **{COORD}**, for guidance and support "
     "throughout the completion of the project work.", 12)
ack3 = P(40)
for r in list(ack3.runs): r._element.getparent().remove(r._element)
runs(ack3, f"We also extend our heartfelt gratitude to the Head of the Department, CSE, **{HOD}**, for continuous encouragement and motivation. It is our "
     "pleasure to acknowledge the support we have received from the institute and various individuals. We would like to thank our Principal, "
     f"**{PRIN}**, D. Y. Patil College of Engineering and Technology, Kolhapur, in particular, for providing constant encouragement, support, and "
     "excellent facilities throughout our academic journey.", 12)
t = T(42)
for i, (n, roll, en) in enumerate(TEAM, 1): set_cell(t.rows[i].cells[1], n, 11)

# ── one-page project report (template 2), placed before the acknowledgement ──
anchor = els[35]


def before(el): anchor.addprevious(el)


def p_before(text="", align=C, size=12, bold=False, after=4):
    p = doc.add_paragraph(); fmt(p, align, 0, after, 1.15); runs(p, text, size, bold); before(p._p); return p


tb = doc.add_table(rows=1, cols=3); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, v in zip(tb.rows[0].cells, ("**Group No:** G45", "**Class:** B.Tech - CSE", "**Division:** D")): set_cell(c, v, 11)
before(tb._tbl)
p_before("", after=2)
for line in ("D. Y. PATIL COLLEGE OF ENGINEERING AND TECHNOLOGY,", "KASABA BAWADA, KOLHAPUR.", "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"):
    p_before(line, size=12, after=0)
p_before("**Academic Year: 2026–2027**", after=4)
p_before(f"**Project Title: {FULL}**", size=13, after=8)
box = doc.add_table(rows=1, cols=1); box.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = box.rows[0].cells[0]; cell.width = Cm(14); cell_borders(cell, "D00000", 12)
box.rows[0].height = Cm(7.2)
set_cell(cell, "[[Insert group photo with guide]]", 13, align=C)
cell.vertical_alignment = 1
before(box._tbl)
p_before(f"({GUIDE} (Guide), " + ", ".join(n for n, _, _ in TEAM) + ")", size=11, after=8)
p = doc.add_paragraph(); fmt(p, J, 0, 0, 1.3)
runs(p, "**Project Description:** " + content.ONE_PAGE, 11.5); before(p._p)
pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE); before(pb._p)

# ── section 2: contents and lists (roman page numbers) ─────────────────────
s1 = doc.sections[0]; page_borders(s1)
sec = doc.add_section(WD_SECTION.NEW_PAGE); page_borders(sec); header_footer(sec, total=False); page_numbers(sec, "lowerRoman", 1)


def title(text, new_page=False):
    p = doc.add_paragraph(); fmt(p, L, 6, 10, 1.15); runs(p, text, 18, True)
    p.paragraph_format.page_break_before = new_page; return p


from docx.enum.style import WD_STYLE_TYPE


def toc_style(name, size, bold, indent, after):
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = bold; st.font.italic = False
    pf = st.paragraph_format; pf.left_indent = Cm(indent); pf.space_before = Pt(0); pf.space_after = Pt(after); pf.line_spacing = 1.15


toc_style("toc 1", 12, True, 0, 3); toc_style("toc 2", 12, False, 0.6, 1); toc_style("toc 3", 11, False, 1.2, 0)
toc_style("table of figures", 12, False, 0, 2)
title("Contents")
p = doc.add_paragraph(); add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the contents.", 12)
title("List of Figures", new_page=True)
p = doc.add_paragraph(); add_field(p, 'TOC \\h \\z \\c "Figure"', "Update field to list figures.", 12)
title("List of Tables", new_page=True)
p = doc.add_paragraph(); add_field(p, 'TOC \\h \\z \\c "Table"', "Update field to list tables.", 12)

# ── section 3: chapters ─────────────────────────────────────────────────────
sec = doc.add_section(WD_SECTION.NEW_PAGE); page_borders(sec); header_footer(sec, total=True); page_numbers(sec, "decimal", 1)


def heading(text, level):
    try:
        st = doc.styles[f"Heading {level}"]
    except KeyError:
        st = None
    p = doc.add_paragraph(style=st) if st is not None else doc.add_paragraph()
    for r in list(p.runs): r._element.getparent().remove(r._element)
    size = {1: 16, 2: 14, 3: 12}[level]
    if level == 1:
        p.paragraph_format.page_break_before = heading.seen
        heading.seen = True
    fmt(p, L, {1: 0, 2: 16, 3: 12}[level], {1: 14, 2: 8, 3: 6}[level], 1.15, keep=True)
    set_font(p.add_run(text), size, True, False, __import__("docx").shared.RGBColor(0, 0, 0))
    ppr = p._p.get_or_add_pPr(); ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level - 1)); ppr.append(ol)
    return p


def para(text, align=J):
    p = doc.add_paragraph(); fmt(p, align, 0, 10, 1.5); runs(p, text, 12)
    p.paragraph_format.widow_control = True; return p


def bullets(items, num=False):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(); fmt(p, J, 0, 6, 1.5, indent=0.9, first=-0.6); p.paragraph_format.widow_control = True
        set_font(p.add_run(f"{i}.\t" if num else "•\t"), 12)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        runs(p, it, 12)


def caption(kind, text):
    p = doc.add_paragraph(); fmt(p, C, 4 if kind == "Figure" else 14, 16 if kind == "Figure" else 6, 1.0, keep=(kind == "Table"))
    set_font(p.add_run("Fig. " if kind == "Figure" else "Tab. "), 11)
    add_field(p, f"SEQ {kind} \\* ARABIC", "1", 11)
    set_font(p.add_run(f": {text}"), 11)


def figure(path, text, width=15.5, max_h=20.0):
    path = pathlib.Path(path)
    w, h = Image.open(path).size
    wcm = min(width, max_h * w / h)
    p = doc.add_paragraph(); fmt(p, C, 12, 0, 1.0, keep=True); p.add_run().add_picture(str(path), width=Cm(wcm))
    caption("Figure", text)


def placeholder_box(text, height=6.0):
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.rows[0].cells[0]; c.width = Cm(14); cell_borders(c, "D00000", 12); tb.rows[0].height = Cm(height); c.vertical_alignment = 1
    set_cell(c, f"[[{text}]]", 12, align=C)


def table(text, header, rows, widths=None, size=10.5):
    caption("Table", text)
    tb = doc.add_table(rows=1 + len(rows), cols=len(header)); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, hcell in enumerate(header):
        set_cell(tb.rows[0].cells[j], f"**{hcell}**", size, align=C); shade(tb.rows[0].cells[j], HEAD_BG)
    # one rule for every table: a column whose values are all short (numbers, IDs, PASS) is centred, text columns are left-aligned
    short = [all(len(str(r[j]).replace("**", "")) <= 12 for r in rows) for j in range(len(header))]
    size = max(size, 10.5)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            set_cell(tb.rows[i].cells[j], v, size, align=C if short[j] else L)
    for row in tb.rows:
        trpr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trpr.append(cs)
        for c in row.cells: c.vertical_alignment = 1
    tblpr = tb._tbl.tblPr; mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 60), ("bottom", 60), ("left", 100), ("right", 100)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblpr.append(mar)
    if widths:
        for row in tb.rows:
            for j, wv in enumerate(widths): row.cells[j].width = Cm(wv)
    trpr = tb.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    for row in tb.rows[:-1]:  # keep a table on one page when it fits
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True
    sp = doc.add_paragraph(); fmt(sp, L, 0, 6, 1.0)
    return tb


def code(title_, text):
    t = para(f"**{title_}**", L); t.paragraph_format.keep_with_next = True; t.paragraph_format.space_before = Pt(8)
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.rows[0].cells[0]; shade(c, "F4F6F8"); cell_borders(c, "9AAAC8", 6); c.text = ""
    first = True
    for line in text.strip("\n").split("\n"):
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        fmt(p, L, 0, 0, 1.0)
        set_font(p.add_run(line if line else " "), 9.5, name="Consolas")
        p.paragraph_format.keep_with_next = True
    trpr = tb.rows[0]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trpr.append(cs)
    sp = doc.add_paragraph(); fmt(sp, L, 0, 8, 1.0)


def page_break():
    """Chapters begin with page_break_before; nothing to insert here."""


heading.seen = False


api = dict(heading=heading, para=para, bullets=bullets, figure=figure, table=table, code=code, page_break=page_break,
           placeholder_box=placeholder_box, caption=caption, S=S, TEAM=TEAM, GUIDE=GUIDE, COORD=COORD)
content.chapters(api)

# the template's trailing empty paragraphs after the acknowledgement table would leave a blank page: remove them
for el in els[43:50]:
    if el.tag.endswith("}p") and not "".join(el.itertext()).strip() and el.getparent() is not None and el.find(qn("w:pPr")) is None or (el.tag.endswith("}p") and not "".join(el.itertext()).strip() and el.find(".//" + qn("w:sectPr")) is None and el.getparent() is not None):
        el.getparent().remove(el)

doc.save(OUT)
print("saved", OUT)

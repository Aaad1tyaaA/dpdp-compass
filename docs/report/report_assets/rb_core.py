"""Low-level helpers for the project-report builder (python-docx on the college template)."""
import copy, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
RED = RGBColor(0xD0, 0x00, 0x00)
HEAD_BG = "DCE3EF"
TITLE = "Multi-Agent RAG System for DPDP Act Compliance Assessment and Citizen Rights Awareness"


def set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color
    return run


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5, keep=False, indent=None, first=None):
    pf = p.paragraph_format
    pf.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = spacing
    if keep: pf.keep_with_next = True
    if indent is not None: pf.left_indent = Cm(indent)
    if first is not None: pf.first_line_indent = Cm(first)
    return p


def runs(p, text, size=12, base_bold=False, italic=False, color=None):
    """Adds text with **bold** and [[placeholder]] markup."""
    for part in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not part: continue
        if part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size, True, italic, color)
        elif part.startswith("[["):
            ph(p, part[2:-2], size)
        else:
            set_font(p.add_run(part), size, base_bold, italic, color)
    return p


def ph(p, text, size=12):
    """Red-circle placeholder: marks data the team still has to supply."""
    set_font(p.add_run("● "), size, True, False, RED, name="Segoe UI Symbol")
    set_font(p.add_run(f"[{text}]"), size, True, False, RED)


def add_field(p, instr, placeholder="1", size=None, italic=None, bold=None):
    def r(tag, **attrs):
        run = p.add_run(); set_font(run, size, bold, italic)
        el = OxmlElement(tag)
        for k, v in attrs.items(): el.set(qn(k), v)
        run._element.append(el); return run
    r("w:fldChar", **{"w:fldCharType": "begin"})
    run = p.add_run(); set_font(run, size, bold, italic)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run._element.append(it)
    r("w:fldChar", **{"w:fldCharType": "separate"})
    set_font(p.add_run(placeholder), size, bold, italic)
    r("w:fldChar", **{"w:fldCharType": "end"})


def border_bottom(p, sz=6, color="000000", side="bottom"):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr"); ppr.append(pbdr)
    b = OxmlElement(f"w:{side}")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "4"), ("w:color", color)): b.set(qn(k), v)
    pbdr.append(b)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcpr.append(s)


def cell_borders(cell, color="000000", sz=8):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "0"), ("w:color", color)): e.set(qn(k), v)
        b.append(e)
    tcpr.append(b)


def set_cell(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    fmt(p, align, 1, 1, 1.0)
    runs(p, str(text), size, bold, color=color)


def page_borders(sect):
    sp = sect._sectPr
    old = sp.find(qn("w:pgBorders"))
    if old is not None: sp.remove(old)
    pb = OxmlElement("w:pgBorders"); pb.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "24"), ("w:color", "000000")): e.set(qn(k), v)
        pb.append(e)
    pgmar = sp.find(qn("w:pgMar"))
    pgmar.addnext(pb)


def page_numbers(sect, fmt_="decimal", start=1):
    sp = sect._sectPr
    pn = sp.find(qn("w:pgNumType"))
    if pn is None:
        pn = OxmlElement("w:pgNumType"); sp.find(qn("w:pgBorders")).addnext(pn) if sp.find(qn("w:pgBorders")) is not None else sp.append(pn)
    for a in list(pn.attrib): del pn.attrib[a]
    pn.set(qn("w:fmt"), fmt_); pn.set(qn("w:start"), str(start))


def header_footer(sect, total=True):
    sect.different_first_page_header_footer = False
    sect.header.is_linked_to_previous = False; sect.footer.is_linked_to_previous = False
    h = sect.header.paragraphs[0]; h.text = ""
    fmt(h, WD_ALIGN_PARAGRAPH.RIGHT, 0, 0, 1.0); set_font(h.add_run(TITLE), 11, False, True); border_bottom(h, 6)
    f = sect.footer.paragraphs[0]; f.text = ""
    fmt(f, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0); border_bottom(f, 6, side="top")
    f.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), alignment=2)
    set_font(f.add_run("Department of Computer Science & Engineering, DYPCET\t"), 11, False, True)
    add_field(f, "PAGE", "1", 11, True)
    if total:
        set_font(f.add_run(" / "), 11, False, True); add_field(f, "SECTIONPAGES", "1", 11, True)


def move_before(el, anchor):
    anchor.addprevious(el)
